from __future__ import annotations

import csv
from pathlib import Path

from nvr_app.services.nvr_scan_service import REPORT_FIELDS


CAMERA_LOOKUP_FIELDS = [
    "摄像头IP",
    "查询结果",
    "NVR地址",
    "NVR名称",
    "通道号",
    "通道名称",
    "序列号",
    "在线状态",
]
NVR_SUMMARY_FIELDS = [
    "NVR地址",
    "NVR名称",
    "状态",
    "物理设备",
    "正常在线",
    "部分不可达",
    "完全离线",
    "状态未知",
    "通道数",
    "错误",
]


def _write_csv(path: str, fieldnames: list[str], rows: list[dict]) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as file_obj:
        writer = csv.DictWriter(file_obj, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    return output


def export_scan_csv(payload: dict, path: str) -> Path:
    return _write_csv(path, REPORT_FIELDS, payload.get("rows", []))


def export_scan_xlsx(payload: dict, path: str) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    device_sheet = workbook.active
    device_sheet.title = "设备结果"
    _write_xlsx_sheet(device_sheet, REPORT_FIELDS, payload.get("rows", []))

    summary_sheet = workbook.create_sheet("NVR汇总")
    _write_xlsx_sheet(
        summary_sheet,
        NVR_SUMMARY_FIELDS,
        payload.get("summaries", []),
    )
    workbook.save(output)
    return output


def camera_lookup_rows(payload: dict) -> list[dict]:
    rows = []
    for result in payload.get("results", []):
        camera_ip = result.get("camera_ip", "")
        matches = result.get("matches", [])
        if not matches:
            result_status = result.get("status", "not_found")
            rows.append(
                {
                    "摄像头IP": camera_ip,
                    "查询结果": (
                        "查询异常"
                        if result_status == "query_error"
                        else "未匹配"
                    ),
                    "NVR地址": "",
                    "NVR名称": "",
                    "通道号": "",
                    "通道名称": "",
                    "序列号": "",
                    "在线状态": "",
                }
            )
            continue

        for match in matches:
            online = match.get("online")
            if online is True:
                online_text = "在线"
            elif online is False:
                online_text = "离线"
            else:
                online_text = "未知"

            rows.append(
                {
                    "摄像头IP": camera_ip,
                    "查询结果": "已匹配",
                    "NVR地址": match.get("nvr_ip", ""),
                    "NVR名称": match.get("nvr_name", ""),
                    "通道号": match.get("channel_id", ""),
                    "通道名称": match.get("channel_name", ""),
                    "序列号": match.get("camera_serial_number", ""),
                    "在线状态": online_text,
                }
            )
    return rows


def export_camera_lookup_csv(payload: dict, path: str) -> Path:
    return _write_csv(path, CAMERA_LOOKUP_FIELDS, camera_lookup_rows(payload))


def export_camera_lookup_xlsx(payload: dict, path: str) -> Path:
    from openpyxl import Workbook

    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    result_sheet = workbook.active
    result_sheet.title = "匹配结果"
    _write_xlsx_sheet(result_sheet, CAMERA_LOOKUP_FIELDS, camera_lookup_rows(payload))

    failure_sheet = workbook.create_sheet("查询异常")
    failure_rows = [
        {
            "NVR地址": item.get("nvr_ip", ""),
            "NVR名称": item.get("nvr_name", ""),
            "错误": item.get("error", ""),
        }
        for item in payload.get("failed_nvrs", [])
    ]
    _write_xlsx_sheet(
        failure_sheet,
        ["NVR地址", "NVR名称", "错误"],
        failure_rows,
    )
    workbook.save(output)
    return output


def _write_xlsx_sheet(sheet, fieldnames: list[str], rows: list[dict]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor="24445C")
    header_font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="微软雅黑", size=10)
    thin = Side(style="thin", color="D8DEE6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    green_fill = PatternFill("solid", fgColor="DDF3E4")
    yellow_fill = PatternFill("solid", fgColor="FFF1CC")
    red_fill = PatternFill("solid", fgColor="F9D7D7")

    for column_index, field in enumerate(fieldnames, start=1):
        cell = sheet.cell(row=1, column=column_index, value=field)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border

    for row_index, row in enumerate(rows, start=2):
        status = str(row.get("状态", row.get("在线状态", "")))
        fill = None
        if status in {"正常在线", "在线", "已匹配"}:
            fill = green_fill
        elif status in {"部分不可达", "状态未知", "未知"}:
            fill = yellow_fill
        elif status in {"网络不可达", "离线", "异常", "查询异常"}:
            fill = red_fill

        for column_index, field in enumerate(fieldnames, start=1):
            cell = sheet.cell(
                row=row_index,
                column=column_index,
                value=row.get(field, ""),
            )
            cell.font = body_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            if fill is not None:
                cell.fill = fill

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, field in enumerate(fieldnames, start=1):
        max_width = len(str(field))
        for row in rows:
            max_width = max(max_width, len(str(row.get(field, ""))))
        sheet.column_dimensions[get_column_letter(column_index)].width = min(
            max(max_width + 3, 10),
            48,
        )

