


#脚本用途：扫描CSV表中添加的监控主机上的监控设备，获取OSD名、IP地址和对应通道。



# -*- coding: utf-8 -*-
import requests
import xml.etree.ElementTree as ET
from requests.auth import HTTPDigestAuth
import csv
import os
import sys
import io
import re

# 渲染依赖
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from nvr_app.services.nvr_scan_service import scan_nvr as scan_nvr_service

# 💡 强行重定向标准输出的编码，防止 Windows 控制台编码错误
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

# ================= 路径配置 =================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CSV_PATH = os.path.join(BASE_DIR, "NVRdevicesTest.csv")

CSV_DIR = os.path.join(BASE_DIR, "reports_csv")
IMG_DIR = os.path.join(BASE_DIR, "reports_imgs")
XLSX_DIR = os.path.join(BASE_DIR, "reports_xlsx")

for d in [CSV_DIR, IMG_DIR, XLSX_DIR]:
    os.makedirs(d, exist_ok=True)

RESULT_CSV_PATH = os.path.join(CSV_DIR, "nvr_temp_report.csv")
RESULT_PNG_PATH = os.path.join(IMG_DIR, "科伦特监控状态一览表.png")
RESULT_XLSX_PATH = os.path.join(XLSX_DIR, "科伦特监控状态一览表.xlsx")

# matplotlib 中文字体
plt.rcParams['font.sans-serif'] = ['SimHei']
plt.rcParams['axes.unicode_minus'] = False

# ================= 工具函数 =================
def _visual_len(s):
    """计算字符串的可视宽度（中文=2，英文=1）"""
    length = 0
    for ch in str(s):
        code = ord(ch)
        if (0x4E00 <= code <= 0x9FFF or 0x3000 <= code <= 0x303F or
                0xFF00 <= code <= 0xFFEF or 0x2E80 <= code <= 0x2EFF):
            length += 2
        else:
            length += 1
    return length

# ================= NVR 设备发现 =================
def get_nvr_list(csv_path):
    nvr_list = []
    if not os.path.exists(csv_path):
        print(f"[-] 错误：找不到 CSV 文件 -> {csv_path}")
        return nvr_list

    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        for row in reader:
            nvr_list.append({k.strip(): v.strip() for k, v in row.items()})
    return nvr_list

def get_base_name(name):
    """提取基础名称（强化对"主/子"、"镜头"等中文后缀的识别）"""
    if not name or name.lower().startswith("camera"):
        return ""
    pattern = r'[-_ ]?(通道|chan|ch|镜头|码流)?\d+$|[-_ ]?(主码流|子码流|主通道|子通道|主|子|主摄|副摄)$'
    base = re.sub(pattern, '', name, flags=re.IGNORECASE).strip()
    return base if len(base) >= 2 else ""

# ================= NVR 巡检核心 =================
def check_single_nvr(nvr):
    """兼容旧 CLI 的扫描入口，实际扫描逻辑由 nvr_scan_service 提供。"""
    ip = nvr.get("ip")
    print(f"\n[SCAN] 正在扫描 NVR: {ip} ...")

    try:
        result = scan_nvr_service(nvr)
    except Exception as exc:
        print(f"  [EXCEPT] 连接失败: {exc}")
        return []

    if result.get("error"):
        print(f"  [!] {result['error']}")

    rows = result.get("rows", [])
    for row in rows:
        status = row["状态"]
        if status == "正常在线":
            prefix = "[正常在线]"
        elif status == "部分不可达":
            prefix = "[部分不可达]"
        else:
            prefix = "[网络不可达]"
        print(
            f"  {prefix} {row['监控设备名称']} | "
            f"IP: {row['IP地址']} | {row['通道属性']} | "
            f"通道: {row['对应/异常通道']}"
        )

    for duplicate in result.get("duplicate_ips", []):
        print(
            f"  [警告] IP {duplicate['ip']} 对应多个设备: "
            f"{', '.join(duplicate['names'])}"
        )

    counts = result.get("counts", {})
    print(
        f"[REPORT] 扫描完成 -> 正常在线: {counts.get('normal', 0)} 台 | "
        f"完全离线: {counts.get('offline', 0)} 台"
    )
    return rows


def _legacy_check_single_nvr(nvr):
    ip = nvr.get('ip')
    user = nvr.get('username')
    pw = nvr.get('password')

    CLR_RESET = "\033[0m"
    CLR_RED = "\033[1;31m"
    CLR_GREEN = "\033[1;32m"
    CLR_YELLOW = "\033[1;33m"
    CLR_CYAN = "\033[1;36m"

    print(f"\n{CLR_CYAN}[SCAN] 正在扫描 NVR: {ip} ...{CLR_RESET}")

    config_url = f"http://{ip}/ISAPI/ContentMgmt/InputProxy/channels"
    status_url = f"http://{ip}/ISAPI/ContentMgmt/InputProxy/channels/status"

    all_channels = {}

    # 1. 获取全量配置
    try:
        config_res = requests.get(config_url, auth=HTTPDigestAuth(user, pw), timeout=8)
        if config_res.status_code == 200:
            config_res.encoding = 'utf-8'
            c_root = ET.fromstring(config_res.text)
            for chan in c_root.findall('.//{*}InputProxyChannel'):
                id_node = chan.find('.//{*}id')
                if id_node is not None and id_node.text:
                    raw_id = id_node.text.strip()

                    name_node = chan.find('.//{*}name')
                    c_name = name_node.text.strip() if (name_node is not None and name_node.text) else f"Camera {raw_id}"

                    c_ip = "无IP"
                    for ip_node in chan.findall('.//{*}ipAddress'):
                        if ip_node is not None and ip_node.text and ip_node.text.strip() not in ["0.0.0.0", ""]:
                            c_ip = ip_node.text.strip()
                            break

                    c_sn = ""
                    for sn_node in chan.findall('.//{*}serialNumber'):
                        if sn_node is not None and sn_node.text:
                            c_sn = sn_node.text.strip()
                            break

                    all_channels[raw_id] = {
                        "id": raw_id,
                        "name": c_name,
                        "ip": c_ip,
                        "sn": c_sn,
                        "online": False
                    }
    except Exception as e:
        print(f"  [!] {CLR_YELLOW}读取配置接口异常: {e}{CLR_RESET}")

    # 2. 获取状态并覆盖合并
    try:
        response = requests.get(status_url, auth=HTTPDigestAuth(user, pw), timeout=8)
        if response.status_code == 200:
            response.encoding = 'utf-8'
            root = ET.fromstring(response.text)
            for channel in root.findall('.//{*}InputProxyChannelStatus'):
                id_node = channel.find('.//{*}id')
                if id_node is not None and id_node.text:
                    raw_id = id_node.text.strip()

                    if raw_id not in all_channels:
                        all_channels[raw_id] = {"id": raw_id, "name": f"Camera {raw_id}", "ip": "无IP", "sn": "", "online": False}

                    online_node = channel.find('.//{*}online')
                    if online_node is not None and online_node.text:
                        is_online = online_node.text.strip().lower() == 'true'
                    else:
                        status_node = channel.find('.//{*}status')
                        status_text = status_node.text.strip().lower() if status_node is not None else ""
                        is_online = (status_text == 'online')

                    all_channels[raw_id]["online"] = is_online

                    ip_node = channel.find('.//{*}ipAddress')
                    if ip_node is not None and ip_node.text and ip_node.text.strip() not in ["0.0.0.0", ""]:
                        all_channels[raw_id]["ip"] = ip_node.text.strip()

                    sn_node = channel.find('.//{*}serialNumber')
                    if sn_node is not None and sn_node.text:
                        all_channels[raw_id]["sn"] = sn_node.text.strip()

                    dev_name_node = channel.find('.//{*}name') or channel.find('.//{*}deviceName')
                    if dev_name_node is not None and dev_name_node.text and all_channels[raw_id]["name"].startswith("Camera"):
                        all_channels[raw_id]["name"] = dev_name_node.text.strip()

    except Exception as e:
        print(f"  {CLR_RED}[EXCEPT] 连接失败: {e}{CLR_RESET}")
        return []

    if not all_channels:
        print(f"  [!] {CLR_YELLOW}未发现任何通道信息{CLR_RESET}")
        return []

    # 3. 建立 IP/SN 共享池
    name_to_ip = {}
    name_to_sn = {}
    for raw_id, info in all_channels.items():
        base = get_base_name(info["name"])
        if base:
            if info["ip"] != "无IP": name_to_ip[base] = info["ip"]
            if info["sn"]: name_to_sn[base] = info["sn"]

    # 4. 字典聚类法
    physical_devices = {}
    for raw_id, info in all_channels.items():
        cam_name = info["name"]
        cam_ip = info["ip"]
        cam_sn = info["sn"]

        base = get_base_name(cam_name)

        if cam_ip == "无IP" and base and base in name_to_ip:
            cam_ip = name_to_ip[base]
        if not cam_sn and base and base in name_to_sn:
            cam_sn = name_to_sn[base]

        if cam_ip != "无IP":
            device_key = f"IP_{cam_ip}"
        elif cam_sn:
            device_key = f"SN_{cam_sn}"
        elif base:
            device_key = f"NAME_{base}"
        else:
            device_key = f"CHAN_{raw_id}"

        if device_key not in physical_devices:
            physical_devices[device_key] = {
                "name": cam_name,
                "ip": cam_ip,
                "channels": {}
            }
        else:
            if physical_devices[device_key]["ip"] == "无IP" and cam_ip != "无IP":
                physical_devices[device_key]["ip"] = cam_ip

        physical_devices[device_key]["channels"][f"D{raw_id}"] = {
            "name": cam_name,
            "online": info["online"]
        }

    # 5. 生成报告
    print(f"\n{CLR_CYAN}[ALERT] [网络状态诊断清单]：{CLR_RESET}")
    print("-" * 85)

    scan_results = []
    online_physical_count = 0
    offline_physical_count = 0

    for dev_key, dev in physical_devices.items():
        chans = dev["channels"]
        total_ch = len(chans)
        channel_type = "单通道" if total_ch == 1 else "双通道"

        offline_ch_ids = [ch_id for ch_id, info in chans.items() if not info["online"]]
        chans_str = ", ".join(chans.keys())

        if len(offline_ch_ids) == total_ch:
            offline_physical_count += 1
            print(f" {CLR_RED}❌ [网络不可达] {dev['name']} | IP: {dev['ip']} | 属性: {channel_type} | 对应通道: {chans_str}{CLR_RESET}")
            scan_results.append({
                "所属NVR": ip,
                "监控设备名称": dev['name'],
                "IP地址": dev['ip'],
                "通道属性": channel_type,
                "对应/异常通道": chans_str,
                "状态": "网络不可达"
            })
        elif len(offline_ch_ids) > 0:
            online_physical_count += 1
            bad_chans_str = ", ".join(offline_ch_ids)
            print(f" {CLR_YELLOW}⚠️ [部分不可达] {dev['name']} | IP: {dev['ip']} | 属性: {channel_type} | 异常通道: {bad_chans_str} 无法连接{CLR_RESET}")
            scan_results.append({
                "所属NVR": ip,
                "监控设备名称": dev['name'],
                "IP地址": dev['ip'],
                "通道属性": channel_type,
                "对应/异常通道": bad_chans_str,
                "状态": "部分不可达"
            })
        else:
            online_physical_count += 1
            print(f" {CLR_GREEN}✅ [正常在线]  {dev['name']} | IP: {dev['ip']} | 属性: {channel_type} | 对应通道: {chans_str}{CLR_RESET}")
            scan_results.append({
                "所属NVR": ip,
                "监控设备名称": dev['name'],
                "IP地址": dev['ip'],
                "通道属性": channel_type,
                "对应/异常通道": chans_str,
                "状态": "正常在线"
            })

    print("-" * 85)
    offline_summary_color = CLR_RED if offline_physical_count > 0 else CLR_RESET
    print(f"[REPORT] 扫描完成 -> {CLR_GREEN}正常在线: {online_physical_count} 台{CLR_RESET} | {offline_summary_color}完全离线: {offline_physical_count} 台{CLR_RESET}")
    print("-" * 85)
    return scan_results

# ================= 报告渲染：PNG 图像 =================
def render_nvr_png(csv_path, png_path):
    """将 NVR CSV 渲染为 PNG 图片，底部附加异常汇总"""
    if not os.path.exists(csv_path):
        return False
    try:
        df = pd.read_csv(csv_path, encoding="utf-8-sig")
        if "所属NVR" in df.columns:
            m = df["所属NVR"] == df["所属NVR"].shift()
            df.loc[m, "所属NVR"] = ""
    except Exception as e:
        print(f"[渲染PNG] 读取 CSV 失败: {e}")
        return False

    headers = list(df.columns)
    row_count = len(df)
    LINE_UNIT_HEIGHT = 0.38
    BASE_ROW_PAD = 0.25

    # ———— 异常汇总数据 ————
    abnormal_mask = df["状态"] != "正常在线"
    abnormal_df = df[abnormal_mask]
    anomaly_lines = []
    if len(abnormal_df) > 0:
        anomaly_lines.append("⚠️ 异常汇总")
        for _, row in abnormal_df.iterrows():
            detail = f"• {row['监控设备名称']} ({row['IP地址']}) — {row['状态']}"
            if row['状态'] == "部分不可达":
                detail += f" [异常通道: {row['对应/异常通道']}]"
            anomaly_lines.append(detail)
        anomaly_lines.append(f"共 {len(abnormal_df)} 台异常设备")
    else:
        anomaly_lines = ["✅ 全部设备正常在线"]

    # 计算总高度：主表格 + 异常汇总区域
    anomaly_text_height = 0.35 * len(anomaly_lines) + 0.6
    fig_width = 17
    fig_height = 1.2 + 0.5 + (row_count * LINE_UNIT_HEIGHT + BASE_ROW_PAD) + anomaly_text_height

    fig, ax = plt.subplots(figsize=(fig_width, fig_height), dpi=200)
    ax.axis('off')

    col_widths_raw = {}
    for col in headers:
        max_w = _visual_len(col)
        for val in df[col].fillna("-"):
            max_w = max(max_w, _visual_len(val))
        col_widths_raw[col] = max_w + 3

    total_w = sum(col_widths_raw.values())
    col_widths = [col_widths_raw[col] / total_w for col in headers]

    tb = ax.table(
        cellText=df.fillna("-").values, colLabels=df.columns,
        cellLoc='center', loc='upper center',
        colColours=['#1F4E79'] * len(headers), colWidths=col_widths
    )

    tb.auto_set_font_size(False)
    tb.set_fontsize(10)

    total_line_units = row_count + 1
    for row_idx in range(row_count + 1):
        for col_idx in range(len(headers)):
            tb[row_idx, col_idx].set_height(1 / total_line_units)

    # 表头样式
    for i in range(len(headers)):
        cell = tb[0, i]
        cell.get_text().set_color('white')
        cell.get_text().set_weight('bold')

    # 数据行着色
    for row_idx in range(1, row_count + 1):
        status_val = str(df.iloc[row_idx - 1, headers.index("状态")]).strip()
        if status_val == "正常在线":
            bg_color, text_color = '#C6EFCE', '#006100'
        elif status_val == "部分不可达":
            bg_color, text_color = '#FFF2CC', '#B25E00'
        else:
            bg_color, text_color = '#FFC6C6', '#9C0006'

        for col_idx in range(len(headers)):
            cell = tb[row_idx, col_idx]
            cell.set_facecolor(bg_color)
            cell.get_text().set_color(text_color)

    # ———— 底部异常汇总文本 ————
    # 计算异常汇总区域的起始 y 坐标（相对于 axes）
    anomaly_y_start = 0.02 + anomaly_text_height / fig_height
    ax.text(0.5, anomaly_y_start * 0.85,
            "\n".join(anomaly_lines),
            transform=ax.transAxes,
            fontsize=10, ha='center', va='bottom',
            fontfamily='sans-serif',
            bbox=dict(boxstyle='round,pad=0.5',
                      facecolor='#FFF3CD' if len(abnormal_df) > 0 else '#D4EDDA',
                      edgecolor='#856404' if len(abnormal_df) > 0 else '#155724',
                      alpha=0.9))

    plt.title("📊 科伦特监控状态一览表", fontsize=15, weight='bold', pad=12, color='#1F4E79')
    plt.savefig(png_path, bbox_inches='tight', dpi=200)
    plt.close()
    return True


# ================= 报告渲染：XLSX（异常汇总在右侧） =================
def render_nvr_xlsx(csv_path, xlsx_path):
    """将 NVR CSV 渲染为 XLSX：主数据在左，异常汇总在右"""
    if not os.path.exists(csv_path):
        return False
    try:
        df = pd.read_csv(csv_path, encoding="utf-8-sig")
    except Exception as e:
        return False

    df = df.fillna("-")
    headers = list(df.columns)

    # ———— 异常数据 ————
    abnormal_mask = df["状态"] != "正常在线"
    abnormal_df = df[abnormal_mask].reset_index(drop=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "监控状态一览"

    FONT_NAME = "微软雅黑"
    HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
    GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
    YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")
    RED_FILL = PatternFill("solid", fgColor="FFC6C6")
    GREEN_TEXT = "006100"
    YELLOW_TEXT = "B25E00"
    RED_TEXT = "9C0006"
    THIN = Side(style="thin", color="D9D9D9")
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    # —— 右侧异常汇总区域列号 ——
    GAP_COL = len(headers) + 1        # 间隔列
    ANOMALY_START_COL = GAP_COL + 1   # 异常汇总起始列
    anomaly_sub_headers = ["⚠️ 异常设备", "IP地址", "异常类型", "异常详情"]
    anomaly_col_count = len(anomaly_sub_headers)
    total_cols = ANOMALY_START_COL + anomaly_col_count - 1

    # ===== 标题行 =====
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    title_cell = ws.cell(row=1, column=1, value="📊 科伦特监控状态一览表")
    title_cell.font = Font(name=FONT_NAME, size=14, bold=True, color="1F4E79")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # ===== 主数据表头 (Row 2) =====
    for c, col_name in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=col_name)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.border = BORDER

    # 间隔列留空但加边框
    gap_cell = ws.cell(row=2, column=GAP_COL, value="")
    gap_cell.border = Border(left=THIN, right=THIN)
    ws.column_dimensions[get_column_letter(GAP_COL)].width = 3

    # ===== 右侧异常汇总表头 (Row 2) =====
    anomaly_header_fill = PatternFill("solid", fgColor="B71C1C")
    for c_offset, sub_h in enumerate(anomaly_sub_headers):
        col_idx = ANOMALY_START_COL + c_offset
        cell = ws.cell(row=2, column=col_idx, value=sub_h)
        cell.font = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
        cell.fill = anomaly_header_fill
        cell.border = BORDER

    # ===== 主数据行 =====
    for r_offset, (_, row) in enumerate(df.iterrows()):
        r = 3 + r_offset
        status_val = str(row.get("状态", "")).strip()
        if status_val == "正常在线":
            bg, text_color = GREEN_FILL, GREEN_TEXT
        elif status_val == "部分不可达":
            bg, text_color = YELLOW_FILL, YELLOW_TEXT
        else:
            bg, text_color = RED_FILL, RED_TEXT

        for c, col_name in enumerate(headers, start=1):
            cell = ws.cell(row=r, column=c, value=str(row[col_name]).strip())
            cell.font = Font(name=FONT_NAME, size=10, color=text_color)
            cell.fill = bg
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 间隔列边框
        ws.cell(row=r, column=GAP_COL).border = Border(left=THIN, right=THIN)

    # ===== 右侧异常汇总数据行 =====
    if len(abnormal_df) > 0:
        for a_idx, (_, arow) in enumerate(abnormal_df.iterrows()):
            r = 3 + a_idx
            status_val = str(arow.get("状态", "")).strip()
            if status_val == "部分不可达":
                anomaly_type = "部分不可达"
                anomaly_detail = f"异常通道: {arow['对应/异常通道']}"
            else:
                anomaly_type = "网络不可达"
                anomaly_detail = "所有通道离线"

            values = [
                str(arow["监控设备名称"]).strip(),
                str(arow["IP地址"]).strip(),
                anomaly_type,
                anomaly_detail
            ]
            for c_offset, val in enumerate(values):
                col_idx = ANOMALY_START_COL + c_offset
                cell = ws.cell(row=r, column=col_idx, value=val)
                cell.font = Font(name=FONT_NAME, size=10, color=RED_TEXT)
                cell.fill = RED_FILL if anomaly_type == "网络不可达" else YELLOW_FILL
                cell.border = BORDER
                cell.alignment = Alignment(horizontal="center", vertical="center")
    else:
        # 无异常时显示绿色对勾
        ws.merge_cells(start_row=3, start_column=ANOMALY_START_COL,
                       end_row=4, end_column=ANOMALY_START_COL + anomaly_col_count - 1)
        ok_cell = ws.cell(row=3, column=ANOMALY_START_COL,
                          value="✅ 全部设备正常在线")
        ok_cell.font = Font(name=FONT_NAME, size=12, bold=True, color=GREEN_TEXT)
        ok_cell.fill = GREEN_FILL
        ok_cell.border = BORDER
        ok_cell.alignment = Alignment(horizontal="center", vertical="center")

    # ===== 列宽自适应 =====
    for c, col_name in enumerate(headers, start=1):
        max_w = max([_visual_len(str(val)) for val in df[col_name]] + [_visual_len(col_name)])
        ws.column_dimensions[get_column_letter(c)].width = max_w + 5

    # 异常汇总列宽
    anomaly_widths = [22, 18, 14, 24]
    for c_offset, w in enumerate(anomaly_widths):
        col_idx = ANOMALY_START_COL + c_offset
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    wb.save(xlsx_path)
    return True


# ================= 检测同一 IP 对应多个设备 =================
def detect_duplicate_ips(all_channels):
    ip_count = {}
    for raw_id, info in all_channels.items():
        cam_ip = info["ip"]
        if cam_ip != "无IP":
            if cam_ip in ip_count:
                ip_count[cam_ip].append(info["name"])
            else:
                ip_count[cam_ip] = [info["name"]]

    for ip, names in ip_count.items():
        if len(names) > 1:
            print(f"[警告] IP {ip} 对应多个设备: {', '.join(names)}")

    return ip_count


# ================= 主入口 =================
if __name__ == "__main__":
    nvr_list = get_nvr_list(CSV_PATH)
    all_results = []
    if nvr_list:
        for nvr in nvr_list:
            results = check_single_nvr(nvr)
            if results:
                all_results.extend(results)

        if all_results:
            # 1. 写 CSV
            fieldnames = ["所属NVR", "监控设备名称", "IP地址", "通道属性", "对应/异常通道", "状态"]
            with open(RESULT_CSV_PATH, "w", encoding="utf-8-sig", newline="") as f:
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                writer.writerows(all_results)
            print(f"✅ 排查完成，数据已保存至: {RESULT_CSV_PATH}")

            # 2. 渲染 PNG 图片
            if render_nvr_png(RESULT_CSV_PATH, RESULT_PNG_PATH):
                print(f"✅ PNG 报告已生成: {RESULT_PNG_PATH}")
            else:
                print(f"⚠️ PNG 渲染失败")

            # 3. 渲染 XLSX 表格（异常汇总在右侧）
            if render_nvr_xlsx(RESULT_CSV_PATH, RESULT_XLSX_PATH):
                print(f"✅ XLSX 报告已生成（含右侧异常汇总）: {RESULT_XLSX_PATH}")
            else:
                print(f"⚠️ XLSX 渲染失败")
        else:
            print("⚠️ 未捕获到任何扫描结果")
    else:
        print("❌ 未找到任何 NVR 设备配置")
