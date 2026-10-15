# -*- coding: utf-8 -*-
import csv
import os
import tempfile
import unittest

from openpyxl import load_workbook

from nvr_app.exporters import (
    camera_lookup_rows,
    export_camera_lookup_xlsx,
    export_scan_csv,
    export_scan_xlsx,
)


class ExporterTests(unittest.TestCase):
    def test_export_scan_csv_omits_credentials(self):
        payload = {
            "rows": [
                {
                    "所属NVR": "192.168.1.10",
                    "监控设备名称": "前门摄像头",
                    "IP地址": "10.0.0.10",
                    "通道属性": "单通道",
                    "对应/异常通道": "D1",
                    "状态": "正常在线",
                    "password": "must-not-export",
                }
            ]
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "scan.csv")
            export_scan_csv(payload, path)
            with open(path, "r", encoding="utf-8-sig", newline="") as file_obj:
                rows = list(csv.DictReader(file_obj))

        self.assertEqual(rows[0]["IP地址"], "10.0.0.10")
        self.assertNotIn("password", rows[0])

    def test_export_scan_xlsx_contains_summary_sheet(self):
        payload = {
            "rows": [],
            "summaries": [
                {
                    "NVR地址": "192.168.1.10",
                    "NVR名称": "",
                    "状态": "正常",
                    "物理设备": 0,
                    "正常在线": 0,
                    "部分不可达": 0,
                    "完全离线": 0,
                    "状态未知": 0,
                    "通道数": 0,
                    "错误": "",
                }
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "scan.xlsx")
            export_scan_xlsx(payload, path)
            workbook = load_workbook(path)

        self.assertEqual(workbook.sheetnames, ["设备结果", "NVR汇总"])

    def test_camera_lookup_export_has_unmatched_row_and_failure_sheet(self):
        payload = {
            "results": [{"camera_ip": "10.0.0.10", "matches": []}],
            "failed_nvrs": [
                {
                    "nvr_ip": "192.168.1.10",
                    "nvr_name": "总部NVR",
                    "error": "timeout",
                }
            ],
        }
        rows = camera_lookup_rows(payload)
        self.assertEqual(rows[0]["查询结果"], "未匹配")

        with tempfile.TemporaryDirectory() as temp_dir:
            path = os.path.join(temp_dir, "lookup.xlsx")
            export_camera_lookup_xlsx(payload, path)
            workbook = load_workbook(path)

        self.assertEqual(workbook.sheetnames, ["匹配结果", "查询异常"])
        self.assertEqual(workbook["查询异常"]["A2"].value, "192.168.1.10")
        self.assertEqual(workbook["查询异常"]["B2"].value, "总部NVR")

    def test_camera_query_error_is_not_reported_as_unmatched(self):
        payload = {
            "results": [
                {
                    "camera_ip": "10.0.0.10",
                    "status": "query_error",
                    "matches": [],
                }
            ],
            "failed_nvrs": [
                {
                    "nvr_ip": "192.168.1.10",
                    "nvr_name": "总部NVR",
                    "error": "ISAPI 请求超时",
                }
            ],
        }

        rows = camera_lookup_rows(payload)

        self.assertEqual(rows[0]["查询结果"], "查询异常")
        self.assertEqual(rows[0]["NVR名称"], "")


if __name__ == "__main__":
    unittest.main()

